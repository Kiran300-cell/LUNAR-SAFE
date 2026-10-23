from io import BytesIO
import json
from datetime import datetime,timezone
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import BarChart,Reference
from openpyxl.formatting.rule import ColorScaleRule
def workbook_bytes(summary,bootstrap,training,manifest):
    wb=Workbook(); dash=wb.active; dash.title="Dashboard"
    dash.append(["LUNAR-SAFE V12 · Evaluation Dashboard"]); dash.append(["Source","Existing held-out test evaluation"])
    dash.append(["Region","Haworth, lunar south pole"]); dash.append(["Test patches per model",600])
    dash.append(["Elevation errors","Relative elevation after patch-centering; vertical units not confirmed"])
    for row in [["Metric","Value"],["Synthetic degradation","5× via non-overlapping 5×5 block averaging"],["Scope","Single region; no independent real 5 m validation"],["Export created UTC",datetime.now(timezone.utc).isoformat()]]: dash.append(row)
    ws=wb.create_sheet("Model Comparison")
    cols=["model","relative_elevation_MAE_m","relative_elevation_RMSE_m","slope_MAE_deg","F1_at_10deg","false_safe_rate_at_10deg"]
    ws.append(cols)
    for _,r in summary[summary.model.isin(["Bicubic","Bilinear","CNN_A","CNN_B"])].iterrows(): ws.append([r.get(c) for c in cols])
    ws.auto_filter.ref=ws.dimensions; ws.freeze_panes="A2"
    chart=BarChart(); chart.title="Relative elevation MAE"; chart.y_axis.title="Relative error"; chart.x_axis.title="Method"
    chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=ws.max_row),titles_from_data=True); chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row)); ws.add_chart(chart,"H2")
    ws.conditional_formatting.add("B2:F20",ColorScaleRule(start_type="min",start_color="63BE7B",mid_type="percentile",mid_value=50,mid_color="FFEB84",end_type="max",end_color="F8696B"))
    for name,data in [("Statistical Validation",bootstrap),("Training Analysis",training)]:
        sh=wb.create_sheet(name)
        if not data.empty:
            sh.append(list(data.columns))
            for row in data.itertuples(index=False,name=None): sh.append([v.item() if hasattr(v,"item") else v for v in row])
            sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
    meta=wb.create_sheet("Dataset Metadata")
    for k,v in manifest.items(): meta.append([k,str(v)])
    instructions=wb.create_sheet("Instructions")
    for row in [["Workbook notes"],["Relative elevation errors are patch-centered. Despite the historic *_m column suffix, vertical units are not confirmed."],["Slope-based threshold classes are exploratory proxies."],["V12 evaluates synthetic 5× degradation from the Haworth DEM."],["Bootstrap intervals resample 36 spatial blocks and depend on the chosen block scale."]]: instructions.append(row)
    for sh in wb.worksheets:
        sh.sheet_view.showGridLines=False
        for cell in sh[1]:
            cell.font=Font(color="FFFFFF",bold=True); cell.fill=PatternFill("solid",fgColor="14243A"); cell.alignment=Alignment(vertical="center")
        for col in sh.columns:
            width=min(42,max(14,max((len(str(c.value or "")) for c in list(col)[:80]),default=10)+2))
            sh.column_dimensions[col[0].column_letter].width=width
    buf=BytesIO(); wb.save(buf); return buf.getvalue()
def pdf_bytes(summary,bootstrap,manifest):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    b=BytesIO(); doc=SimpleDocTemplate(b,pagesize=letter); styles=getSampleStyleSheet(); story=[]
    story += [Paragraph("LUNAR-SAFE | V12 Terrain Reconstruction Evaluation",styles["Title"]),Spacer(1,12),
      Paragraph("Research prototype report · existing test evaluation; no newly executed model results.",styles["BodyText"]),Spacer(1,12),
      Paragraph("Method",styles["Heading2"]),Paragraph("Synthetic 5× degradation was generated using non-overlapping 5×5 block averaging on spatially separated Haworth patches. CNN checkpoints were selected using validation elevation L1. Relative elevation metrics are computed after patch-centering.",styles["BodyText"]),Spacer(1,10)]
    cols=["Method","Rel. MAE","Rel. RMSE","Slope MAE°","F1 @10°","False-safe"]
    rows=[cols]
    for _,r in summary[summary.model.isin(["Bicubic","Bilinear","CNN_A","CNN_B"])].iterrows():
        rows.append([str(r.model)]+[f"{float(r[k]):.5f}" for k in ["relative_elevation_MAE_m","relative_elevation_RMSE_m","slope_MAE_deg","F1_at_10deg","false_safe_rate_at_10deg"]])
    t=Table(rows,repeatRows=1); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#14243A")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.4,colors.grey),("FONTSIZE",(0,0),(-1,-1),8)])); story += [Paragraph("Test metrics (patch means)",styles["Heading2"]),t,Spacer(1,12)]
    if not bootstrap.empty:
        story.append(Paragraph(f"Paired spatial block bootstrap comparisons are provided for {len(bootstrap)} metric contrasts against bilinear interpolation; they use 36 spatial blocks and should be interpreted conditional on block scale.",styles["BodyText"]))
    story += [Spacer(1,12),Paragraph("Limitations",styles["Heading2"]),Paragraph("Synthetic degradation only; one lunar region; independent real 5 m validation unavailable; source DEM artifacts may affect reconstruction; vertical units are not confirmed. Slope threshold labels are hazard proxies, not certified landing-safety labels. Higher spatial detail does not establish real-world landing safety.",styles["BodyText"])]
    doc.build(story); return b.getvalue()

def workbook_bytes(summary,bootstrap,training,manifest,pooled=None,patches=None):
    from openpyxl.utils import get_column_letter
    wb=Workbook(); dash=wb.active; dash.title="Dashboard"
    dash.append(["LUNAR-SAFE V12 · Evaluation Dashboard"])
    for row in [
        ["Evaluation type","Existing held-out test evaluation"],
        ["Region","Haworth, lunar south pole"],
        ["Test patches per model",600],
        ["Degradation","Synthetic 5× using non-overlapping 5×5 block averaging"],
        ["Elevation metric interpretation","Relative, patch-centered; vertical units are not confirmed"],
        ["Hazard labels","Exploratory slope-threshold proxies; not safety certification"]]: dash.append(row)
    dash.append(["Method","Relative MAE","Relative RMSE","Slope MAE (deg)","F1 @ 10°","False-safe @ 10°"])
    selected=summary[summary.model.isin(["Bicubic","Bilinear","CNN_A","CNN_B"])] if "model" in summary else summary
    for _,r in selected.iterrows():
        dash.append([r.get("model")]+[r.get(k) for k in ["relative_elevation_MAE_m","relative_elevation_RMSE_m","slope_MAE_deg","F1_at_10deg","false_safe_rate_at_10deg"]])
    def add_frame(title,frame):
        sh=wb.create_sheet(title)
        if frame is None or frame.empty:
            sh.append(["No source data available"]); return sh
        sh.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):
            sh.append([v.item() if hasattr(v,"item") else v for v in row])
        sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
        return sh
    comp=add_frame("Model Comparison",selected[["model","relative_elevation_MAE_m","relative_elevation_RMSE_m","slope_MAE_deg","F1_at_10deg","false_safe_rate_at_10deg"]])
    if comp.max_row>1:
        chart=BarChart(); chart.title="Relative elevation MAE"; chart.y_axis.title="Relative error"; chart.x_axis.title="Method"
        chart.add_data(Reference(comp,min_col=2,min_row=1,max_row=comp.max_row),titles_from_data=True)
        chart.set_categories(Reference(comp,min_col=1,min_row=2,max_row=comp.max_row)); comp.add_chart(chart,"H2")
    add_frame("Hazard Analysis",pooled)
    add_frame("Statistical Validation",bootstrap)
    add_frame("Training Analysis",training)
    meta=wb.create_sheet("Dataset Metadata")
    for key,value in (manifest or {}).items(): meta.append([key,json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else str(value)])
    add_frame("Patch-Level Results",patches)
    instructions=wb.create_sheet("Instructions")
    for row in [["Workbook notes"],["All values are from supplied V12 output files; this workbook does not report a new model run."],["Elevation errors are relative and patch-centered; vertical units are unconfirmed."],["Spatial block-bootstrap intervals use 36 blocks and depend on block scale."],["The synthetic 5× test does not validate independent real 5 m inputs."],["Slope threshold labels are exploratory hazard proxies, not landing-safety certification."]]: instructions.append(row)
    for sh in wb.worksheets:
        sh.sheet_view.showGridLines=False
        sh.freeze_panes=sh.freeze_panes or "A2"
        sh.auto_filter.ref=sh.auto_filter.ref or (sh.dimensions if sh.max_row>1 else None)
        for cell in sh[1]:
            cell.font=Font(color="FFFFFF",bold=True); cell.fill=PatternFill("solid",fgColor="14243A"); cell.alignment=Alignment(vertical="center")
        for col in sh.columns:
            width=min(44,max(13,max((len(str(c.value or "")) for c in list(col)[:100]),default=10)+2))
            sh.column_dimensions[get_column_letter(col[0].column)].width=width
    buf=BytesIO(); wb.save(buf); return buf.getvalue()



def pdf_bytes(summary,bootstrap,manifest,pooled=None,training=None,config=None):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from config import SOURCE_URL
    b=BytesIO(); doc=SimpleDocTemplate(b,pagesize=letter,rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=40)
    styles=getSampleStyleSheet(); story=[]
    story += [Paragraph("LUNAR-SAFE",styles["Title"]),Paragraph("V12 · Terrain Reconstruction & Hazard Proxy Evaluation",styles["Heading2"]),
      Paragraph("Research prototype report · existing test evaluation; no new model run was performed.",styles["BodyText"]),Spacer(1,12)]
    story += [Paragraph("Executive summary",styles["Heading2"]),
      Paragraph("The supplied V12 results compare bicubic and bilinear interpolation with CNN-A and CNN-B on spatially separated Haworth test patches. CNN-A has the lowest patch-mean relative elevation MAE and RMSE in this evaluation. This does not establish performance on real 5 m data or other lunar regions.",styles["BodyText"]),Spacer(1,10),
      Paragraph("Methodology",styles["Heading2"]),
      Paragraph("The pipeline uses non-overlapping 5×5 block averaging to synthesize nominal 5× degradation. The 12,060 × 11,660 Haworth DEM has nominal 1 m horizontal sampling and a lunar polar stereographic CRS. There are 600 patches each in train, validation and test splits, separated into spatial column bands. CNN-A optimizes elevation L1; CNN-B adds first-order gradient L1 with weight 0.10. Both select checkpoints using validation elevation L1. Test elevation errors are calculated after subtracting the LR-derived patch mean.",styles["BodyText"]),Spacer(1,10)]
    cols=["Method","Rel. MAE","Rel. RMSE","Slope MAE°","F1 @10°","False-safe"]
    rows=[cols]
    for _,r in summary[summary.model.isin(["Bicubic","Bilinear","CNN_A","CNN_B"])].iterrows():
        rows.append([str(r.model)]+[f"{float(r[k]):.5f}" for k in ["relative_elevation_MAE_m","relative_elevation_RMSE_m","slope_MAE_deg","F1_at_10deg","false_safe_rate_at_10deg"]])
    def table(data,widths=None):
        t=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#14243A")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#9AAAC3")),("FONTSIZE",(0,0),(-1,-1),7.5),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#EEF3F8")])]))
        return t
    story += [Paragraph("Model comparison · patch means",styles["Heading2"]),table(rows),Spacer(1,9),
      Paragraph("Hazard analysis",styles["Heading2"]),
      Paragraph("Slope threshold classes are exploratory proxies. The original output reports per-patch means and pooled pixel-count statistics; these estimands are distinct. False-safe rate describes reference hazardous cells predicted below threshold. It is not a mission safety probability.",styles["BodyText"]),Spacer(1,8)]
    if pooled is not None and not pooled.empty:
        p=pooled[pooled.threshold_deg==10.0] if "threshold_deg" in pooled else pooled
        cols=[c for c in ["model","threshold_deg","TP","FP","FN","TN","pooled_F1","pooled_false_safe_rate"] if c in p]
        if cols:
            story.append(table([cols]+[[str(v) for v in row] for row in p[cols].head(10).itertuples(index=False,name=None)]))
    story += [Spacer(1,10),Paragraph("Statistical validation",styles["Heading2"])]
    if bootstrap is not None and not bootstrap.empty:
        bdf=bootstrap[bootstrap.candidate.isin(["CNN_A","CNN_B"])]
        cols=[c for c in ["candidate","metric","difference_candidate_minus_baseline","ci_2_5","ci_97_5","n_spatial_blocks"] if c in bdf]
        story.append(table([cols]+[[str(v) for v in row] for row in bdf[cols].itertuples(index=False,name=None)]))
        story.append(Paragraph("Paired spatial block-bootstrap intervals use 36 blocks and 2,000 replicates. Their interpretation depends on block scale; intervals do not establish cross-region generalization.",styles["BodyText"]))
    story += [Spacer(1,8),Paragraph("Reproducibility and source",styles["Heading2"])]
    cfg=config or {}
    story.append(Paragraph(f"Scale factor: {cfg.get('scale_factor','5')} · patch size: {cfg.get('hr_patch_size','100')} px · seeds: {cfg.get('seeds',[42,143,244])} · bootstrap seed: {cfg.get('bootstrap_seed','20261023')}. Official source: {SOURCE_URL}",styles["BodyText"]))
    story += [Spacer(1,8),Paragraph("Limitations",styles["Heading2"]),
      Paragraph("Synthetic degradation only; one region; no independent real 5 m validation; source DEM artifacts may affect reconstruction and evaluation. Vertical band units are not confirmed, so historic *_m metric suffixes must not be read as verified absolute metre errors. Slope threshold labels are hazard proxies, not certified lunar landing safety labels. Higher DEM detail does not demonstrate improved real-world landing safety.",styles["BodyText"])]
    doc.build(story); return b.getvalue()
